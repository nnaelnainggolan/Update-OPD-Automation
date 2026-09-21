"""Check daily workbook/sheet mapping before a long batch; never writes Excel."""
import io
import json
from pathlib import Path
from openpyxl import load_workbook
from excel_report import MONTHS


def check_report(projects, root, target_date):
    options = json.loads((Path(root) / 'excel_config.json').read_text(encoding='utf-8'))
    if not options.get('enabled', False):
        return {'status': 'disabled'}
    report = Path(options['drive_report_folder']) / f'{target_date.day:02d} {MONTHS[target_date.month]} {target_date.year}.xlsx'
    if not report.is_file():
        raise FileNotFoundError(f'Siapkan file laporan tanggal tujuan sebelum menjalankan batch: {report}')
    if report.with_name('~$' + report.name).exists():
        raise PermissionError('Tutup file laporan di Excel sebelum menjalankan batch.')
    book = load_workbook(io.BytesIO(report.read_bytes()), read_only=True)
    try:
        names = [s.strip() for s in book.sheetnames]
        missing = [p['report_sheet'] for p in projects
                   if names.count(p['report_sheet'].strip()) > 1
                   or (names.count(p['report_sheet'].strip()) == 0 and not p.get('create_missing_sheet', False))]
        if missing:
            raise ValueError('Sheet tidak ditemukan tunggal pada laporan: ' + ', '.join(missing))
        return {'status': 'ready', 'file': str(report), 'projects': len(projects),
                'sheets_to_create': [p['report_sheet'] for p in projects if p['report_sheet'].strip() not in names]}
    finally:
        book.close()
