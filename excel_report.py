"""Write verified capture results to an existing daily workbook, with backup."""
from pathlib import Path
from datetime import date, datetime
import io
import hashlib
import json
import logging
import os
import shutil
import tempfile
import zipfile

from openpyxl import load_workbook
from openpyxl.drawing.image import Image as ExcelImage
from PIL import Image

MONTHS = ['', 'JANUARI','FEBRUARI','MARET','APRIL','MEI','JUNI','JULI','AGUSTUS','SEPTEMBER','OKTOBER','NOVEMBER','DESEMBER']
SLOTS = {'BROADBAND': ('A2', 1, 16), 'METRO': ('A18', 17, 33)}


def digest(path):
    h = hashlib.sha256()
    with open(path, 'rb') as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b''):
            h.update(chunk)
    return h.hexdigest()


def occupied(sheet, first, last):
    for image in sheet._images:
        anchor = image.anchor
        marker = getattr(anchor, '_from', None)
        if marker is None:
            return True  # Unknown placement: do not risk overlapping an existing image.
        start = marker.row + 1
        end_marker = getattr(anchor, 'to', None)
        end = end_marker.row + 1 if end_marker else start + max(1, int(image.height / 20))
        if start <= last and end >= first:
            return True
    return False


def save_reports(results, capture_dir, cfg, root):
    root, capture_dir = Path(root), Path(capture_dir)
    options = json.loads((root / 'excel_config.json').read_text(encoding='utf-8'))
    if not options.get('enabled', False):
        return {'status':'disabled'}
    good = [r for r in results if r.get('status') in {'graph_present','no_data'}]
    if not good:
        return {'status':'no_valid_capture'}
    dates = {date.fromisoformat(r['date']) for r in good}
    if len(dates) != 1 or len({r['network'] for r in good}) != len(good):
        raise ValueError('Tanggal campuran atau capture jaringan ganda; Excel tidak diubah.')
    target_date = next(iter(dates))
    folder = Path(options['drive_report_folder'])
    report = folder / f'{target_date.day:02d} {MONTHS[target_date.month]} {target_date.year}.xlsx'
    if not report.is_file():
        raise FileNotFoundError(f'File tanggal tujuan belum tersedia: {report}')
    if report.with_name('~$' + report.name).exists():
        raise PermissionError('Tutup file laporan di Excel sebelum menjalankan capture.')
    lock = report.with_name(report.name + '.opd.lock')
    fd = os.open(lock, os.O_CREAT | os.O_EXCL | os.O_WRONLY)
    os.close(fd)
    temporary = None
    workbook = None
    transfer = None
    stage = "membaca sumber"
    try:
        logging.info('EXCEL v2.7.1: membaca salinan sumber ke memori')
        source_bytes = report.read_bytes()
        baseline = hashlib.sha256(source_bytes).hexdigest()
        workbook = load_workbook(io.BytesIO(source_bytes))
        stage = 'menambahkan gambar' 
        matches = [s for s in workbook.worksheets if s.title.strip() == cfg['report_sheet'].strip()]
        if not matches and cfg.get('create_missing_sheet', False):
            title = cfg['report_sheet'].strip()
            if not title or len(title) > 31 or any(ch in title for ch in '[]:*?/\\'):
                raise ValueError('Nama sheet baru tidak valid: ' + title)
            sheet = workbook.create_sheet(title)
            sheet['A17'] = 'Link Metro'
            if 'BROADBAND' in cfg.get('networks', ['METRO','BROADBAND']):
                sheet['A1'] = 'Link Broadband'
            sheet.column_dimensions['A'].width = 24
            for row in range(1, 34):
                sheet.row_dimensions[row].height = 20
            logging.info('Sheet baru disiapkan: %s', title)
        elif len(matches) == 1:
            sheet = matches[0]
        else:
            raise ValueError('Sheet tujuan tidak ditemukan tunggal: ' + cfg['report_sheet'])
        outcomes = []
        for result in good:
            network = result['network']
            anchor, first, last = SLOTS[network]
            if occupied(sheet, first, last):
                logging.warning('Dilewati: %s!%s sudah berisi gambar.', sheet.title, anchor)
                outcomes.append({'network':network,'cell':anchor,'status':'skipped_existing'})
                continue
            image_path = (capture_dir / result['graph']).resolve()
            if image_path.parent != capture_dir.resolve():
                raise ValueError('Lokasi gambar di luar folder capture.')
            with Image.open(image_path) as im:
                width, height = im.size
                im.verify()
            # Fit Broadband before row 17, preserving all row heights and labels.
            start_row = 2 if network == 'BROADBAND' else 18
            available_px = sum((sheet.row_dimensions[r].height or sheet.sheet_format.defaultRowHeight or 15) * 96/72
                               for r in range(start_row, start_row + 15)) - 8
            scale = min(1, 1100/width, min(330, available_px)/height)
            picture = ExcelImage(str(image_path))
            picture.width, picture.height = max(1,round(width*scale)), max(1,round(height*scale))
            sheet.add_image(picture, anchor)
            # In-memory anchors are strings; next slot should not treat them as unknown.
            from openpyxl.drawing.spreadsheet_drawing import _check_anchor
            picture.anchor = _check_anchor(picture)
            outcomes.append({'network':network,'cell':anchor,'status':'inserted'})
        if not any(r['status'] == 'inserted' for r in outcomes):
            return {'status':'skipped_existing','file':str(report),'items':outcomes}
        backup_dir = root / 'backups' / target_date.isoformat()
        backup_dir.mkdir(parents=True, exist_ok=True)
        backup = backup_dir / (report.stem + '_' + datetime.now().strftime('%H%M%S_%f') + '.xlsx')
        if digest(report) != baseline:
            raise RuntimeError('File tujuan berubah selama proses; penyimpanan dibatalkan.')
        stage = "membuat backup lokal"
        backup.write_bytes(source_bytes)
        if digest(backup) != baseline:
            raise RuntimeError('Backup tidak cocok dengan file awal; penyimpanan dibatalkan.')
        handle, name = tempfile.mkstemp(prefix='opd_save_', suffix='.xlsx')
        os.close(handle)
        temporary = Path(name)
        stage = "menyimpan Excel pada disk lokal"
        logging.info("EXCEL: %s", stage)
        workbook.save(temporary)
        stage = "memverifikasi arsip lokal"
        with zipfile.ZipFile(temporary) as archive:
            if archive.testzip():
                raise ValueError('Verifikasi arsip Excel gagal.')
        if digest(report) != baseline:
            raise RuntimeError('File tujuan berubah sebelum penyimpanan; hasil tidak ditimpa.')
        stage = 'menyalin hasil terverifikasi ke Drive lokal'
        logging.info('EXCEL: %s', stage)
        fd, transfer_name = tempfile.mkstemp(prefix='opd_transfer_', suffix='.xlsx', dir=folder)
        os.close(fd)
        transfer = Path(transfer_name)
        shutil.copyfile(temporary, transfer)
        if digest(transfer) != digest(temporary):
            raise RuntimeError('Salinan Drive belum cocok dengan hasil lokal; file tujuan tidak diganti.')
        if digest(report) != baseline:
            raise RuntimeError('File tujuan berubah saat transfer; file tujuan tidak diganti.')
        stage = 'mengganti file tujuan'
        os.replace(transfer, report)
        logging.info('Excel tersimpan: %s. Sinkronisasi online mengikuti Google Drive for Desktop.', report)
        return {'status':'saved','file':str(report),'sheet':sheet.title,'backup':str(backup),'items':outcomes}
    except Exception as exc:
        logging.exception('EXCEL gagal pada tahap: %s', stage)
        raise RuntimeError(f'{stage}: {exc}') from exc
    finally:
        if transfer:
            transfer.unlink(missing_ok=True)
        if workbook:
            workbook.close()
        if temporary:
            temporary.unlink(missing_ok=True)
        lock.unlink(missing_ok=True)
